#!/usr/bin/env python3
"""Verify release links, editable/printable content, archives and page geometry."""
from pathlib import Path
from html.parser import HTMLParser
from urllib.parse import urlsplit, unquote
import hashlib
import json
import re
import zipfile
import pymupdf as fitz
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DIST = ROOT / 'dist'
SITE = ROOT / 'release'
VERSION = (ROOT / 'VERSION').read_text().strip()


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self.scripts = []
    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in ['a', 'link', 'img']:
            self.links.append(attrs.get('href', attrs.get('src', '')))
        if tag in ['script', 'form', 'iframe']:
            self.scripts.append(tag)


def main():
    pages = sorted(SITE.rglob('*.html'))
    count = 0
    for p in pages:
        parser = Links()
        parser.feed(p.read_text())
        assert not parser.scripts, (p, parser.scripts)
        for link in parser.links:
            url = urlsplit(link)
            if url.scheme or url.netloc or not url.path or url.path == '/':
                continue
            dest = (p.parent / unquote(url.path)).resolve()
            assert dest.is_relative_to(SITE.resolve()), (p, link)
            assert dest.exists(), (p, link)
            count += 1
    pdfs = sorted(DIST.glob(f'*-v{VERSION}.pdf'))
    for p in pdfs:
        d = fitz.open(p)
        for page in d:
            assert abs(page.rect.width - 595.3) < 2 and abs(page.rect.height - 841.9) < 2
            assert 'v' + VERSION in page.get_text(), (p, page.number)
            for block in page.get_text('dict')['blocks']:
                for line in block.get('lines', []):
                    for span in line['spans']:
                        x0,y0,x1,y1 = span['bbox']
                        assert x0 >= 10 and x1 <= page.rect.width - 10 and y0 >= 5 and y1 <= page.rect.height - 5, (p, page.number, span['text'])
        if p.stem.startswith('printable-starter'):
            assert len(d) == 11
    # Check that source paragraphs and cells survive conversion to PDF.
    normalise = lambda s: re.sub(r'\W+', '', s).lower()
    for p in DIST.glob(f'*-v{VERSION}.docx'):
        doc = Document(p)
        with fitz.open(p.with_suffix('.pdf')) as pdf:
            flattened = normalise(' '.join(page.get_text() for page in pdf))
        texts = [para.text for para in doc.paragraphs]
        texts += [para.text for table in doc.tables for row in table.rows for cell in row.cells for para in cell.paragraphs]
        for text in texts:
            needle = normalise(text)
            if len(needle) > 10:
                assert needle in flattened, (p.name, text)
    for p in DIST.glob(f'*-v{VERSION}.zip'):
        with zipfile.ZipFile(p) as z:
            assert z.testzip() is None
            assert any(n.endswith('/LICENSE') for n in z.namelist())
            assert not any(n.startswith('/') or '..' in Path(n).parts or any(x in n for x in ['.env', 'receipt.json', '.eml', '/outreach/', '/.git/']) for n in z.namelist())
    for line in (DIST / 'SHA256SUMS').read_text().splitlines():
        digest, name = line.split('  ', 1)
        assert hashlib.sha256((DIST / name).read_bytes()).hexdigest() == digest
    print(json.dumps({'status':'passed','html_pages':len(pages),'local_links_checked':count,'pdf_files':len(pdfs),'docx_files':len(list(DIST.glob(f'*-v{VERSION}.docx'))),'checks':['local links','no scripts/forms/iframes','A4 geometry','page text within bounds','versions on every PDF page','DOCX text preserved in PDF','archive integrity and scope','download checksums']},indent=2))


if __name__ == '__main__':
    main()
