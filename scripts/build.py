#!/usr/bin/env python3
"""Build blank DOCX/PDF templates and static release bundles. No patient input."""
from pathlib import Path
import hashlib
import html
import json
import logging
import re
import shutil
import subprocess
import tempfile
import zipfile

import pymupdf as fitz
import markdown
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'dist'
VERSION = (ROOT / 'VERSION').read_text().strip()
DATE = '22 September 2026'
LOG = logging.getLogger('continuity-build')


def text(p, value, bold=False, size=None):
    r = p.add_run(value)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    return r


def base(title, subtitle='Your notes, in your own words', size=11):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin, s.bottom_margin = Mm(14), Mm(16)
    s.left_margin, s.right_margin = Mm(17), Mm(17)
    s.header_distance, s.footer_distance = Mm(6), Mm(7)
    for name in ['Normal', 'Heading 1', 'Heading 2', 'Title']:
        st = d.styles[name]
        st.font.name = 'Liberation Sans'
        st.font.color.rgb = RGBColor.from_string('182B32')
        fonts = st.element.get_or_add_rPr().rFonts
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():
                del fonts.attrib[attr]
        borders = st.element.xpath('./w:pPr/w:pBdr')
        for border in borders:
            border.getparent().remove(border)
    normal = d.styles['Normal']
    normal.font.size = Pt(size)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, pts in [('Heading 1', 22), ('Heading 2', 12), ('Title', 23)]:
        st = d.styles[name]
        st.font.size = Pt(pts)
        st.paragraph_format.space_before = Pt(9 if name == 'Heading 2' else 0)
        st.paragraph_format.space_after = Pt(5)
    h = s.header.paragraphs[0]
    text(h, 'PATIENT-HELD CONTINUITY PACK', size=8)
    foot = s.footer.paragraphs[0]
    foot.paragraph_format.space_after = Pt(0)
    text(foot, f'v{VERSION} | {DATE} | Patient support, not a clinical record | ', size=8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    foot._p.append(field)
    notice = s.footer.add_paragraph()
    notice.paragraph_format.space_after = Pt(0)
    text(notice, '© 2026 Andrew Brogan | MIT licence: continuity.stationarystore.ie/LICENSE.txt', size=8)
    d.core_properties.author = 'Andrew Brogan'
    d.core_properties.title = title
    d.core_properties.subject = f'Patient-held Continuity Pack v{VERSION}: blank public resource'
    d.core_properties.keywords = f'patient support, blank template, v{VERSION}'
    d.add_heading(title, 0)
    if subtitle:
        p = d.add_paragraph()
        text(p, subtitle, size=10)
    return d


def box(d, label, height, prompt=''):
    t = d.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    t.autofit = False
    t.columns[0].width = Mm(176)
    row = t.rows[0]
    row.height = Mm(height)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    p = row.cells[0].paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    text(p, label, bold=True, size=10.5)
    if prompt:
        for line in prompt.split('\n'):
            p = row.cells[0].add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            if '\n' in prompt:
                p.paragraph_format.line_spacing = Pt(16)
            text(p, line, size=9)
    pr = row._tr.get_or_add_trPr()
    pr.append(OxmlElement('w:cantSplit'))
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1
    text(p, '', size=2)
    p.paragraph_format.line_spacing = Pt(3)


def grid(d, headers, widths, rows=5, height=23):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Mm(width)
    for cell, label, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Mm(width)
        text(cell.paragraphs[0], label, bold=True, size=9)
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for _ in range(rows):
        r = t.add_row()
        r.height, r.height_rule = Mm(height), WD_ROW_HEIGHT_RULE.AT_LEAST
        r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, width in zip(r.cells, widths):
            cell.width = Mm(width)
    return t


def save(d, name):
    path = OUT / f'{name}-v{VERSION}.docx'
    d.save(path)
    LOG.info('Wrote %s', path.name)


def forms():
    d = base('Appointment sheet', 'One sheet per visit. Blank boxes are fine. Continue on the back if needed.')
    box(d, 'Name or chosen identifier:                         Date / time:', 13, 'Service / clinician:                                      Visit type: in person / phone / video')
    box(d, 'BEFORE | What matters most to me today?', 22)
    box(d, 'My questions | Choose up to three', 27, '1.\n2.\n3.')
    box(d, 'Changes since the last visit / information to bring', 19, 'Bring your current medicines list and relevant letters. Mark uncertain details “not sure”.')
    box(d, 'DURING | What I understood from the conversation', 34, '“Can I read back what I understood and check whether I have missed anything?”')
    box(d, 'NEXT | What happens, who arranges it, and by when?', 30, 'For tests: who reviews the result, how will I hear, and whom do I contact if I do not?')
    box(d, 'What changes mean I should seek help sooner, and how?', 23, 'Record advice from the clinician. This sheet does not provide medical advice.')
    box(d, 'Still unclear / contact for questions', 17)
    p = d.add_paragraph('Checking understanding is not sign-off. No clinician signature is needed. Ask about unclear instructions before changing treatment. Do not wait for paperwork when you need care.')
    p.paragraph_format.space_after = Pt(0)
    for r in p.runs:
        r.font.size = Pt(9)
    save(d, 'appointment-sheet')

    d = base('Current information', 'Your summary as of a date. It may be incomplete. Keep source documents alongside it.')
    box(d, 'Name / chosen identifier:                           Updated on:', 15, 'Prepared by (if someone helped):                         With my permission: yes / no')
    box(d, 'What matters to me / communication and access needs', 28)
    box(d, 'Current health information I want to share', 40, 'Include the source and date where known. Mark questions or uncertainty clearly.')
    box(d, 'Allergies or previous reactions I know about', 29, 'What caused the reaction? What happened? Source / date? Write “not sure” if uncertain.')
    box(d, 'Medicines information is kept here:', 19, 'List date / source:                           Includes non-prescription products? yes / no / not sure')
    box(d, 'People / services involved in my care', 30, 'Name or service | role | contact route (keep personal details inside the folder)')
    box(d, 'Current plan and things waiting to happen', 27, 'Source and date | See Waiting and follow-up for open items.')
    box(d, 'What needs checking / information that does not match', 23)
    save(d, 'current-information')

    d = base('My medicines information', 'Optional worksheet. You can use an existing current pharmacy or HSE medicines list instead.')
    box(d, 'Name / chosen identifier:                           Updated on:', 16, 'Information came from / list date:')
    d.add_paragraph('Include medicines you take regularly or occasionally, non-prescription products, vitamins and supplements. Copy instructions from a current source; write “not sure” and ask a pharmacist or clinician about gaps or differences. This is not a prescription.')
    grid(d, ['Name and strength', 'How much I take, how I take it and when', 'Regular / as needed; instructions or limits', 'Source / date; notes or questions'], [43, 47, 43, 43], rows=5, height=25)
    box(d, 'Allergies / reactions: what caused them and what happened?', 26, 'Include source / date. Write “not sure” if uncertain.')
    d.add_paragraph('Do not change treatment because of this worksheet. Keep it current after advice from your clinical team and mark replaced lists “superseded”. Use another dated sheet if you need more rows.')
    d.add_page_break()
    d.add_heading('My pharmacy labels', 0)
    p = d.add_paragraph('Side 2 of 2. Print the medicines sheet double-sided, flipping on the long edge.')
    for r in p.runs: r.font.size = Pt(9)
    p = d.add_paragraph('Ask your pharmacy whether it can supply a spare dispensing label, or attach a copy. Keep the original label and instructions on the medicine packaging. Labels may become outdated; keep the list on the front current and ask about any differences.')
    for r in p.runs: r.font.size = Pt(10)
    d.add_paragraph('Name / chosen identifier:                         Sheet date:')
    for number in range(1, 4):
        box(d, f'Label {number} | Added on:                         Matches medicine / row on front:', 56,
            'Status: current / changed / stopped / not sure       Checked on:')
    p = d.add_paragraph('Attach a spare label or copy in each space. If it is replaced, mark the old label “superseded” and date the change. Use another copy of this side if needed. A label is a reference, not confirmation that instructions are still current.')
    for r in p.runs: r.font.size = Pt(9)
    save(d, 'medicines-list')

    d = base('Waiting and follow-up', 'An optional reminder for you. This sheet does not monitor results or contact anyone.')
    box(d, 'Name / chosen identifier:                           Sheet started on:', 15)
    d.add_paragraph('Write the agreed contact route and expected date. If no date was given, ask. Do not assume that no news means a normal result. Seek care when you need it; do not wait for this list.')
    grid(d, ['What am I waiting for?', 'Who arranges / reviews it? Contact route', 'Expected by / when to ask', 'What happened? Date / next step / closed'], [44, 48, 36, 48], rows=6, height=28)
    d.add_paragraph('Close an item only when you know what happened or have an agreed next step. If a date changes, record the new date and who gave it. Keep the original appointment sheet or letter.')
    save(d, 'follow-up-tracker')

    dividers = [
        ('Current information', 'Start here. Your dated summary and current medicines information.', 'Keep source documents alongside your notes. Mark uncertainty. Move replaced summaries to Older information.'),
        ('Appointments', 'One sheet for each visit. Put the next blank sheet at the front.', 'Before: your priorities and questions. During: what you understood. After: the next step, owner, date and contact route.'),
        ('Letters and care plans', 'Copies of letters and plans, newest first.', 'Keep originals intact. Note where each document came from. If two documents disagree, ask the relevant service to clarify.'),
        ('Tests and results', 'Copies of test information and results, newest first.', 'Keep the accompanying explanation or contact route. Put missing results in Waiting and follow-up. Do not interpret an unexplained result using this pack.'),
        ('Waiting and follow-up', 'Open items: what happens next, who arranges it, and by when.', 'Use the tracker if it helps. Do not assume silence means a normal result. The folder does not chase or monitor anything for you.'),
        ('Older information', 'Superseded summaries and older papers you want to keep.', 'Label replaced summaries “superseded” with a date. Add older letters gradually. An incomplete archive does not stop you starting today.'),
    ]
    d = base('Folder dividers', 'Print A4, single-sided. One divider per section. Colour is not needed.')
    for i, (title, desc, tip) in enumerate(dividers):
        if i:
            d.add_page_break()
        d.add_paragraph('\n\n')
        p = d.add_paragraph()
        text(p, f'{i+1:02}', bold=True, size=48)
        d.add_heading(title, 0)
        d.add_paragraph(desc)
        d.add_paragraph(tip)
        d.add_paragraph('\nUse only the sections that help you. Keep personal information inside your folder, not on its outside cover.')
    save(d, 'folder-dividers')

    d = base('Pilot information and consent', 'Proposed anonymous usability activity | Participant keeps this sheet', size=10.5)
    box(d, 'Host / facilitator:                                      Written contact:', 15, 'Facilitator fills this in and explains the activity before inviting consent.')
    d.add_paragraph('We are testing whether the blank pack is understandable and usable. This is not treatment or a test of health outcomes. Start with a fictional example; using it at a routine visit is optional. You may use a supporter, skip questions or stop. Taking part or declining does not affect your care.')
    d.add_paragraph('Keep every completed health sheet yourself. Do not share names, diagnoses, services, appointment dates or identifying stories in feedback. No health records, audio/video or names are requested. The facilitator records only a random code, consent date and design feedback, stored privately and locally. Free text can accidentally identify someone: keep comments about the design.')
    d.add_paragraph('Choose a random code without initials or date of birth. There is no code-to-name list. Keep the code to request removal of your feedback until 14 days after your feedback session. After that, responses may be combined in an anonymous summary and cannot be individually removed. Individual responses are deleted within 30 days of that summary.')
    d.add_paragraph('Possible burdens include time, writing effort or discomfort. You do not need to explain your health or complete every task. Ask the facilitator about the activity; ask your usual service about care. No benefit is guaranteed.')
    d.add_heading('My choice', 2)
    for statement in [
        'I understand the purpose and have had a chance to ask questions.',
        'I know participation is optional and I can stop without affecting care.',
        'I understand what feedback is kept, the removal window and the privacy limits.',
        'I agree to give design feedback under these arrangements.',
    ]:
        d.add_paragraph('[  ] ' + statement)
    box(d, 'My random code:                                      Consent date:', 14)
    box(d, 'Feedback session date:                      Removal request deadline:', 14, 'Facilitator completes the deadline as feedback date + 14 days; participant keeps a copy.')
    d.add_paragraph('No signature required for this proposed approach. Do not begin if the contact route or privacy arrangements are unclear. The host must resolve any additional governance requirements before recruitment.')
    save(d, 'pilot-consent')

    d = base('Pilot feedback', 'Design feedback only. No names, diagnoses, health records or identifying stories.', size=10.5)
    box(d, 'Random participant code:                              Feedback date:', 14)
    d.add_paragraph('Tick one answer per row. Skip anything you do not want to answer. “Not tried” is a useful answer.')
    grid(d, ['Could you...', 'Yes', 'With help', 'No', 'Not tried'], [112, 14, 20, 14, 16], rows=0)
    t = d.tables[-1]
    for label in ['Find where current information belongs?', 'Write one question for an appointment?', 'Find where to note the next step and contact?', 'Explain that a clinician need not sign the sheet?']:
        cells = t.add_row().cells
        cells[0].text = label
        for c in cells[1:]:
            c.text = '[  ]'
    box(d, 'What helped, if anything?', 26)
    box(d, 'What was confusing, tiring or unnecessary?', 26)
    box(d, 'Any difficulty reading, writing, handling or keeping it private?', 26)
    box(d, 'What would you change or remove?', 26)
    d.add_paragraph('Would you choose to use it?  [  ] Yes  [  ] Maybe  [  ] No  [  ] Prefer not to say')
    d.add_paragraph('Did you try it at a visit?  [  ] Yes  [  ] No  [  ] Prefer not to say\nDo not add the service, appointment date or any clinical details.')
    d.add_paragraph('Keep your code if you may want your response removed within the period on your consent sheet. Give this feedback privately to the agreed facilitator, not to a public website.')
    save(d, 'pilot-feedback')


def markdown_docs():
    for path in sorted((ROOT / 'content').glob('*.md')):
        lines = path.read_text().splitlines()
        title = lines[0].removeprefix('# ')
        compact = path.stem in ['advocacy-brief', 'quick-start', 'how-to-use-appointment-sheet']
        d = base(title, '', size=12 if path.stem in ['quick-start', 'how-to-use-appointment-sheet'] else 11)
        for block in '\n'.join(lines[1:]).strip().split('\n\n'):
            block = block.strip()
            if not block:
                continue
            if block.startswith('Version '):
                continue
            if block.startswith('## '):
                d.add_heading(block[3:], 2)
            else:
                d.add_paragraph(block.replace('**', ''))
        save(d, path.stem)


def build_pdfs():
    with tempfile.TemporaryDirectory(prefix='continuity-lo-') as profile:
        cmd = ['libreoffice', f'-env:UserInstallation={Path(profile).as_uri()}', '--headless', '--convert-to', 'pdf', '--outdir', str(OUT)]
        result = subprocess.run(cmd + [str(p) for p in sorted(OUT.glob(f'*-v{VERSION}.docx'))], capture_output=True, text=True, timeout=150)
        if result.returncode:
            raise RuntimeError(f'LibreOffice conversion failed: {result.stdout} {result.stderr}')
        if result.stderr.strip():
            LOG.warning('LibreOffice: %s', result.stderr.strip())
    expected = {name: 1 for name in ['appointment-sheet', 'current-information', 'medicines-list', 'follow-up-tracker', 'quick-start', 'how-to-use-appointment-sheet', 'advocacy-brief', 'pilot-consent', 'pilot-feedback']}
    expected['folder-dividers'] = 6
    expected['medicines-list'] = 2
    report = {}
    for docx in sorted(OUT.glob(f'*-v{VERSION}.docx')):
        p = docx.with_suffix('.pdf')
        if not p.is_file():
            raise RuntimeError(f'Missing PDF: {p}')
        doc = fitz.open(p)
        name = p.stem.removesuffix(f'-v{VERSION}')
        for page in doc:
            if f'v{VERSION}' not in page.get_text():
                raise RuntimeError(f'Missing version on {p.name} page {page.number+1}')
        if name in expected and len(doc) != expected[name]:
            raise RuntimeError(f'{name}: expected {expected[name]} pages, got {len(doc)}')
        report[p.name] = {'pages': len(doc), 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()}
        LOG.info('Checked %s: %s pages', p.name, len(doc))
    bundle = fitz.open()
    for name in ['quick-start', 'current-information', 'medicines-list', 'appointment-sheet', 'follow-up-tracker', 'folder-dividers']:
        with fitz.open(OUT / f'{name}-v{VERSION}.pdf') as d:
            bundle.insert_pdf(d)
    bundle.save(OUT / f'printable-starter-pack-v{VERSION}.pdf', garbage=4, deflate=True)
    (OUT / 'build-report.json').write_text(json.dumps(report, indent=2) + '\n')


def previews():
    """Build actual-page previews and static, downloadable carousel content."""
    target = ROOT / 'site' / 'media' / 'pages'
    target.mkdir(parents=True, exist_ok=True)
    descriptions = {
        'appointment-sheet': ('Appointment sheet', 'One sheet for one visit: your priorities, questions, understanding and next steps.'),
        'current-information': ('Current information', 'A dated summary of what matters now, with sources and room for uncertainty.'),
        'medicines-list': ('Medicines information', 'Your current list on the front. Three spaces for spare pharmacy labels or copies on the reverse.'),
        'follow-up-tracker': ('Waiting and follow-up', 'Keep the next step, who arranges it and when to ask in view.'),
        'folder-dividers': ('Folder dividers', 'Six printable dividers. Use the sections that help you.'),
        'quick-start': ('Quick start', 'Begin with today’s information and the next appointment. Older letters can wait.'),
        'how-to-use-appointment-sheet': ('Using the appointment sheet', 'A short explanation of what to write before, during and after a visit.'),
        'advocacy-brief': ('Advocacy brief', 'A one-page introduction to the problem, the idea and the proposed review.'),
        'evidence': ('Evidence and limitations', 'The sources behind the design, and what has not yet been established.'),
        'advocacy-targets': ('Advocacy contacts', 'Verified groups to approach about review and patient support.'),
        'talk-script': ('15-minute talk', 'A timed script for explaining the folder and inviting feedback.'),
        'pilot-plan': ('Proposed pilot', 'A small, voluntary trial with consent and questions for feedback.'),
        'pilot-consent': ('Pilot consent', 'Explain the proposed pilot and record a person’s choice to take part.'),
        'pilot-feedback': ('Pilot feedback', 'Find out what helped, what did not and what should change.'),
    }
    slides = []
    for slug, (title, description) in descriptions.items():
        doc = fitz.open(OUT / f'{slug}-v{VERSION}.pdf')
        for i, page in enumerate(doc):
            filename = f'{slug}-v{VERSION}-{i+1}.webp'
            pix = page.get_pixmap(matrix=fitz.Matrix(840 / page.rect.width, 840 / page.rect.width))
            pix.pil_save(target / filename, format='WEBP', quality=86)
            label = 'Pharmacy labels' if slug == 'medicines-list' and i == 1 else title
            if slug == 'folder-dividers':
                label = ['Current information', 'Appointments', 'Letters and care plans', 'Tests and results', 'Waiting and follow-up', 'Older information'][i] + ' divider'
            idx = len(slides) + 1
            loading = 'eager' if idx == 1 else 'lazy'
            url = f'downloads/{slug}-v{VERSION}.pdf#page={i+1}'
            slides.append(f'''<article class="preview-page" data-title="{html.escape(label, quote=True)}" role="group" aria-roledescription="slide" aria-label="{idx}: {html.escape(label, quote=True)}">
<a class="preview-image" href="{url}" aria-label="Open {html.escape(label, quote=True)} PDF, page {i+1}"><img src="media/pages/{filename}" width="840" height="1188" loading="{loading}" decoding="async" alt="{html.escape(label, quote=True)}, page {i+1} of {len(doc)}. Preview of the blank printable page."></a>
<div class="preview-detail"><p class="eyebrow">{'For your folder' if idx <= 13 else 'For advocates'} · v{VERSION} · Page {i+1} of {len(doc)}</p><h3>{html.escape(label)}</h3><p>{html.escape(description)}</p><p class="preview-links"><a href="{url}">Open PDF ↗</a><a href="downloads/{slug}-v{VERSION}.docx">Editable Word ↓</a></p></div></article>''')
    page = ROOT / 'site' / 'downloads.html'
    raw = page.read_text()
    raw = re.sub(r'<!-- PREVIEW PAGES START -->.*?<!-- PREVIEW PAGES END -->', '<!-- PREVIEW PAGES START -->\n' + '\n'.join(slides) + '\n<!-- PREVIEW PAGES END -->', raw, flags=re.S)
    raw = re.sub(r'(data-preview-count[^>]*>)\d+ / \d+', rf'\g<1>1 / {len(slides)}', raw)
    raw = re.sub(r'(id="preview-slider"[^>]*max=")\d+', rf'\g<1>{len(slides)}', raw)
    page.write_text(raw)
    LOG.info('Rendered %s actual-page previews', len(slides))


def bundles():
    def put(z, p):
        z.write(p, f'patient-held-continuity-pack-v{VERSION}/{p.relative_to(ROOT)}')
    with zipfile.ZipFile(OUT / f'patient-held-continuity-pack-v{VERSION}.zip', 'w', zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.iterdir()):
            if p.suffix in ['.pdf', '.docx'] and f'-v{VERSION}.' in p.name:
                z.write(p, f'patient-held-continuity-pack-v{VERSION}/downloads/{p.name}')
        for p in sorted((ROOT / 'site').rglob('*')):
            if p.is_file():
                target = f'patient-held-continuity-pack-v{VERSION}/{p.relative_to(ROOT / "site")}'
                if p.suffix == '.html':
                    offline = p.read_text().replace('href="/"', 'href="https://stationarystore.ie/"')
                    offline = re.sub(r'href="(downloads/[^\"]+(?:\.zip|SHA256SUMS))"', r'href="https://continuity.stationarystore.ie/\1"', offline)
                    z.writestr(target, offline)
                else:
                    z.write(p, target)
        for folder in ['content', 'docs']:
            for p in sorted((ROOT / folder).rglob('*')):
                if p.is_file():
                    put(z, p)
        for name in ['README.md', 'LICENSE', 'VERSION', 'CHANGELOG.md']:
            put(z, ROOT / name)
    with zipfile.ZipFile(OUT / f'patient-held-continuity-pack-source-v{VERSION}.zip', 'w', zipfile.ZIP_DEFLATED) as z:
        for folder in ['content', 'docs', 'scripts', 'site']:
            for p in sorted((ROOT / folder).rglob('*')):
                if p.is_file() and '__pycache__' not in p.parts:
                    put(z, p)
        for name in ['README.md', 'LICENSE', 'VERSION', 'CHANGELOG.md', '.gitignore', 'GUIDE.md', 'requirements.txt']:
            put(z, ROOT / name)
    files = sorted(p for p in OUT.iterdir() if p.is_file() and p.name != 'SHA256SUMS')
    (OUT / 'SHA256SUMS').write_text(''.join(f'{hashlib.sha256(p.read_bytes()).hexdigest()}  {p.name}\n' for p in files))


def guides():
    target = ROOT / 'site' / 'guides'
    target.mkdir(exist_ok=True)
    inputs = {p.stem: p for p in (ROOT / 'content').glob('*.md')}
    inputs.update({'readme': ROOT / 'README.md', 'digital-roadmap': ROOT / 'docs/digital-roadmap.md', 'contributing': ROOT / 'docs/CONTRIBUTING.md'})
    for slug, p in inputs.items():
        raw = p.read_text()
        # Autolink bare source URLs, leaving ordinary Markdown links alone.
        raw = re.sub(r'(?<![\(<])(https://[^\s<>]+)', r'<\1>', raw)
        body = markdown.markdown(raw, extensions=['fenced_code', 'tables'])
        title = html.escape(p.read_text().splitlines()[0].lstrip('# '))
        page = f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="referrer" content="no-referrer"><title>{title} | Patient-held Continuity Pack</title><link rel="stylesheet" href="../style.css?rev=20260922c"></head><body><a class="skip" href="#main">Skip to content</a><main id="main" class="guide"><a class="back" href="../index.html">Back to the pack and downloads</a>{body}</main><footer>Patient-held Continuity Pack · v{VERSION} · {DATE} · <a href="../LICENSE.txt">MIT licence</a></footer></body></html>'''
        (target / f'{slug}.html').write_text(page)
    shutil.copy2(ROOT / 'LICENSE', ROOT / 'site' / 'LICENSE.txt')


def assemble_site():
    dest = ROOT / 'release'
    shutil.copytree(ROOT / 'site', dest, dirs_exist_ok=True)
    downloads = dest / 'downloads'
    downloads.mkdir(exist_ok=True)
    for p in OUT.iterdir():
        if p.is_file():
            shutil.copy2(p, downloads / p.name)


def main():
    logging.basicConfig(level=logging.INFO, format='%(levelname)s %(message)s')
    OUT.mkdir(exist_ok=True)
    forms()
    markdown_docs()
    build_pdfs()
    guides()
    previews()
    bundles()
    assemble_site()
    LOG.info('Release v%s complete: %s', VERSION, OUT)


if __name__ == '__main__':
    main()
